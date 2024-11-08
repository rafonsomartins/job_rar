import time
import pyautogui
import pyperclip
import pandas as pd
import pygetwindow as gw
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import getpass
from datetime import datetime

def wait_for_window(window_title, timeout=7):
    start_time = time.time()
    while time.time() - start_time < timeout:
        active_window_title = gw.getActiveWindow().title
        if active_window_title == window_title:
            return True
        time.sleep(0.5)
    return False

def open_sap(username, password):
    pyautogui.press('winleft')
    time.sleep(0.2)
    if not wait_for_window("Search"):
        return print("Couldn't open Windows menu")
    pyautogui.typewrite("SAP")
    time.sleep(0.2)
    pyautogui.press('enter')

    if not wait_for_window("SAP Logon 740"):
        return print("Couldn't open SAP Logon 740")

    pyautogui.press('enter')

    if not wait_for_window("SAP"):
        return print("Couldn't open RAR(1)/000 SAP")

    pyautogui.typewrite(str(username))
    time.sleep(0.1)
    pyautogui.press('down')
    pyautogui.typewrite(str(password))
    pyautogui.press('enter')

    if not wait_for_window("SAP Easy Access"):
        return print("Couldn't open SAP Easy Access")

    pyautogui.hotkey("alt", "space")
    pyautogui.press("x")

    time.sleep(0.2)

    return 1

def close_sap():
    pyautogui.hotkey("alt", "f4")
    time.sleep(0.5)
    pyautogui.press("tab")
    pyautogui.press("enter")
    return

def save_from_clipboard(start_x, start_y, end_x, end_y, type):
    pyperclip.copy('0')

    pyautogui.moveTo(start_x, start_y)
    pyautogui.mouseDown()
    pyautogui.moveTo(end_x, end_y)
    pyautogui.mouseUp()

    pyautogui.hotkey("ctrl", "c")
    time.sleep(0.1)
    text = pyperclip.paste()
    if type == "i":
        return float(text.replace('.', '').replace(',', '.'))
    if type == "t":
        return text

def open_roteiro():
    coordinates = [
        (300, 250),  # y1
        (300, 270),  # y2
        (300, 288),  # y3
        (300, 310),  # y4
        (300, 331)   # y5
    ]

    dates = []
    for x, y in coordinates:
        date_str = save_from_clipboard(1444, y, 1373, y, "t")
        if date_str != '0':
            dates.append((datetime.strptime(date_str, '%d.%m.%Y'), x, y))
        else:
            break

    if not dates:
        return 0

    max_date, max_x, max_y = max(dates)

    pyautogui.doubleClick(max_x, max_y)
    if not wait_for_window("Exibir RotProdRep: síntese de operação", 2):
        return 0

    return 1

def make_dic(material, atividade):
    x1 = save_from_clipboard(160, 470, 25, 466, "t")
    x2 = save_from_clipboard(157, 492, 24, 489, "t")
    if x2 != '0':   
        x3 = save_from_clipboard(160, 515, 24, 512, "t")
        if x3 != '0':
            x4 = save_from_clipboard(166, 533, 23, 533, "t")
            if x4 != '0':
                x5 = save_from_clipboard(166, 554, 23, 554, "t")
                if x5 != '0':
                    x6 = save_from_clipboard(166, 578, 23, 578, "t")
                    if x6 != '0':
                        dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t"), x2: save_from_clipboard(248, 489, 168, 492, "i"), f"UM {x2}": save_from_clipboard(280, 492, 254, 492, "t"), x3: save_from_clipboard(248, 512, 166, 514, "i"), f"UM {x3}": save_from_clipboard(280, 513, 254, 513, "t"), x4: save_from_clipboard(250, 536, 169, 532, "i"), f"UM {x4}": save_from_clipboard(280, 534, 254, 534, "t"), x5: save_from_clipboard(250, 556, 169, 556, "i"), f"UM {x5}": save_from_clipboard(280, 555, 254, 555, "t"), x6: save_from_clipboard(250, 578, 169, 578, "i"), f"UM {x6}": save_from_clipboard(280, 578, 254, 578, "t")}
                    else:
                        dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t"), x2: save_from_clipboard(248, 489, 168, 492, "i"), f"UM {x2}": save_from_clipboard(280, 492, 254, 492, "t"), x3: save_from_clipboard(248, 512, 166, 514, "i"), f"UM {x3}": save_from_clipboard(280, 513, 254, 513, "t"), x4: save_from_clipboard(250, 536, 169, 532, "i"), f"UM {x4}": save_from_clipboard(280, 534, 254, 534, "t"), x5: save_from_clipboard(250, 556, 169, 556, "i"), f"UM {x5}": save_from_clipboard(280, 555, 254, 555, "t")}
                else:
                    dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t"), x2: save_from_clipboard(248, 489, 168, 492, "i"), f"UM {x2}": save_from_clipboard(280, 492, 254, 492, "t"), x3: save_from_clipboard(248, 512, 166, 514, "i"), f"UM {x3}": save_from_clipboard(280, 513, 254, 513, "t"), x4: save_from_clipboard(250, 536, 169, 532, "i"), f"UM {x4}": save_from_clipboard(280, 534, 254, 534, "t")}
            else:
                dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t"), x2: save_from_clipboard(248, 489, 168, 492, "i"), f"UM {x2}": save_from_clipboard(280, 492, 254, 492, "t"), x3: save_from_clipboard(248, 512, 166, 514, "i"), f"UM {x3}": save_from_clipboard(280, 513, 254, 513, "t")}
        else:
            dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t"), x2: save_from_clipboard(248, 489, 168, 492, "i"), f"UM {x2}": save_from_clipboard(280, 492, 254, 492, "t")}
    else:
        dic = {"Material": material, "Atividade": atividade, "Qtd. básica": int(save_from_clipboard(267, 380, 171, 378, "t").replace('.', '')), x1: save_from_clipboard(248, 466, 169, 468, "i"), f"UM {x1}": save_from_clipboard(280, 466, 254, 467, "t")}
    return dic

def open_atividade(material):
    result = []
    for y in [273, 288, 311, 332, 353, 376, 395, 410, 434, 458, 479, 500, 520, 540, 561]:
        atividade = save_from_clipboard(607, y, 369, y, "t")
        if atividade != '0':
            pyautogui.doubleClick(x=607, y=y)
            if not wait_for_window("Exibir RotProdRep: detalhe de operação"):
                result.append({"Material": "Erro", "Erro": f"detalhe de operação, atividade: {atividade}"})
                pyautogui.click(x=266, y=51, clicks=4, interval=0.5)
                wait_for_window("Exibir RotProdRep: síntese de roteiros")
                continue
            result.append(make_dic(material, atividade))
            pyautogui.click(x=266, y=51)
            wait_for_window("Exibir RotProdRep: síntese de operação")
            time.sleep(0.3)
        else:
            while not wait_for_window("Exibir RotProdRep: 1ª tela", 0.5):
                pyautogui.click(x=266, y=51)
                break
    return result

def run_ca23(material):
    time.sleep(0.3)
    pyautogui.typewrite(str(material))
    pyautogui.press('down')
    pyautogui.typewrite("ra01")
    pyautogui.press('enter')
    time.sleep(2.5)
    pyautogui.moveTo(256, 179)
    pyautogui.mouseDown()
    pyautogui.moveTo(254, 120)
    pyautogui.mouseUp()
    time.sleep(0.3)
    pyautogui.doubleClick(233, 224)
    if not wait_for_window("Exibir RotProdRep: síntese de roteiros", 0.6):
        return open_atividade(material)
    if not open_roteiro():
        return
    time.sleep(0.3)
    return open_atividade(material)

def main():
    username = input("SAP username: ")
    password = getpass.getpass("SAP password: ")

    input_file = str(input("Path para o ficheiro com os materiais[1 para default]: "))
    if input_file == '1':
        input_file = 'result.csv'

    try:
        input_df = pd.read_csv(input_file, delimiter=';', skip_blank_lines=True)
    except:
        return (print(f"Couldn't read {input_df}"))
    materials = input_df['Material'].astype(int)

    output_file = str(input("Nome do ficheiro final [1 para default]: "))
    if output_file == '1':
        output_file = '1_roteiros'
    try:
        existing_df = pd.read_excel(f'{output_file}.xlsx')
    except FileNotFoundError:
        existing_df = pd.DataFrame()

    result = []

    if open_sap(username, password):
        pyautogui.typewrite("CA23")
        pyautogui.press('enter')
        if wait_for_window("Exibir RotProdRep: 1ª tela"):
            for material in materials:
                buffer = run_ca23(material)
                result.append(buffer)
                time.sleep(0.1)
        else:
            return print("Couldn't open CA23")
        close_sap()
        flat_result = [item for sublist in result for item in sublist]
        if existing_df.empty:
            df = pd.DataFrame(flat_result)
        else:
            df = pd.concat([existing_df, pd.DataFrame(flat_result)], ignore_index=True)

        df.to_csv(f'{output_file}.csv', index=False, sep=';', decimal=',', encoding='utf-8-sig')
        sheet = 'Roteiros'
        df.to_excel(f'{output_file}.xlsx', index=False, sheet_name=sheet)
        wb = openpyxl.load_workbook(filename = f'{output_file}.xlsx')
        tab = Table(displayName="df", ref=f'A1:{openpyxl.utils.get_column_letter(df.shape[1])}{len(df)+1}')
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
        wb[sheet].add_table(tab)
        for column in wb[sheet].columns:
            max_length = 0
            column = [cell for cell in column if cell.value is not None]
            if column:
                max_length = max(len(str(cell.value)) for cell in column)
            adjusted_width = (max_length + 2) * 1.2
            wb[sheet].column_dimensions[column[0].column_letter].width = adjusted_width
        wb.save(f'{output_file}.xlsx')
    else:
        return print("Couldn't open SAP")

if __name__ == "__main__":
    start = time.time()
    main()
    print(f"\nProgram duration: {round(time.time()-start)} seconds")
